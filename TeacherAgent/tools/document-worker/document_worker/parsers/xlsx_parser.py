"""XLSX 文档解析器（基于 openpyxl）。"""

from __future__ import annotations

from ..schema import PageOrSheet, ParseResult

MAX_FILE_SIZE = 50 * 1024 * 1024  # 50MB（预览解析上限）


def parse_xlsx(file_path: str) -> ParseResult:
    """解析 XLSX 文件，提取文本内容。不执行宏。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ParseResult(ok=False, error="openpyxl 未安装")

    try:
        import os
        file_size = os.path.getsize(file_path)
        if file_size > MAX_FILE_SIZE:
            return ParseResult(
                ok=False,
                error=f"文件大小超过当前上限 50MB（{file_size // (1024 * 1024)}MB）。建议压缩或拆分后重试。"
            )

        # data_only=True 不计算公式，只读取缓存值
        wb = load_workbook(file_path, read_only=True, data_only=True)

        sheets: list[PageOrSheet] = []
        all_text_parts: list[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append("\t".join(cells))
            sheet_text = "\n".join(rows)
            sheets.append(PageOrSheet(name=sheet_name, text=sheet_text))
            all_text_parts.append(sheet_text)

        wb.close()

        plain_text = "\n\n".join(all_text_parts)
        title = os.path.basename(file_path)

        warnings: list[str] = []
        if not plain_text.strip():
            warnings.append("表格未提取到数据")

        return ParseResult(
            ok=True,
            file_type="xlsx",
            title=title,
            pages_or_sheets=sheets,
            plain_text=plain_text,
            warnings=warnings,
        )
    except Exception as e:
        return ParseResult(ok=False, error=f"XLSX 解析失败: {e}")
